"""Generate paper.docx from paper.md for arXiv submission via Word -> PDF."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

md = Path("docs/paper/FALSIFIABILITY_GATED_DISCOVERY.md").read_text(encoding="utf-8")
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)

lines = md.split("\n")
in_code = False
table_rows: list[list[str]] = []

def flush_table():
    global table_rows
    if not table_rows:
        return
    cols = max(len(r) for r in table_rows)
    t = doc.add_table(rows=len(table_rows), cols=cols)
    t.style = "Table Grid"
    for ri, row in enumerate(table_rows):
        for ci, cell in enumerate(row):
            if ci < cols:
                t.rows[ri].cells[ci].text = cell.strip()
    doc.add_paragraph()
    table_rows = []

for line in lines:
    if line.startswith("**Draft"):
        continue
    if line.startswith("# ") and not line.startswith("## "):
        flush_table()
        p = doc.add_heading(line[2:].strip(), level=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith("### "):
        flush_table()
        doc.add_heading(line[4:].strip(), level=2)
    elif line.startswith("## "):
        flush_table()
        doc.add_heading(line[3:].strip(), level=1)
    elif line.startswith("---"):
        flush_table()
    elif line.startswith("|"):
        cells = [c.strip() for c in line.split("|")[1:-1]]
        if cells and not all(set(c) <= set("-: ") for c in cells):
            table_rows.append(cells)
    elif line.startswith("```"):
        flush_table()
        in_code = not in_code
    elif in_code:
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
    elif line.startswith("- "):
        flush_table()
        doc.add_paragraph(line[2:].strip(), style="List Bullet")
    elif line.strip():
        flush_table()
        doc.add_paragraph(line.strip())

flush_table()
out = Path(r"C:\Users\云丘\Desktop\投稿文件2\paper.docx")
doc.save(str(out))
print(f"OK: {out.name} ({out.stat().st_size:,} bytes)")
